from pathlib import Path
import tkinter as tk

import pandas as pd
import pytest
from openpyxl import Workbook

from ms_feature_db_matcher.config import DEFAULT_DNA_PATH, DEFAULT_OIL_ADDUCT_PATH, DEFAULT_RNA_PATH
from ms_feature_db_matcher.gui import (
    AppState,
    DEFAULT_RNA_SUBTYPE_MODE,
    MatcherApp,
    database_badge_text,
    describe_mode,
    mode_label,
    path_badge_text,
    rna_subtype_enabled,
    run_matching,
    status_appearance,
)
from ms_feature_db_matcher.matcher import MatchMode, RnaSubtypeMode


@pytest.fixture(scope="module")
def tk_root():
    root = tk.Tk()
    root.withdraw()
    try:
        yield root
    finally:
        root.destroy()


@pytest.fixture
def matcher_app(tk_root):
    tk_root.withdraw()
    tk_root.geometry("640x690")
    for child in tk_root.winfo_children():
        child.destroy()

    app = MatcherApp(tk_root, AppState.create())
    tk_root.update_idletasks()
    try:
        yield app, tk_root
    finally:
        for child in tk_root.winfo_children():
            child.destroy()
        tk_root.withdraw()


def test_app_state_preloads_default_paths_and_output_dir() -> None:
    state = AppState.create()

    assert state.output_dir.exists()
    assert "datatables.xlsx" in str(state.dna_db_path)
    assert "natural_modifications.xlsx" in str(state.rna_db_path)
    assert DEFAULT_OIL_ADDUCT_PATH.name == "oil_adduct_database.xlsx"
    assert DEFAULT_OIL_ADDUCT_PATH.exists()


def test_describe_mode_explains_dna_plus_rna_ordering() -> None:
    summary = describe_mode(MatchMode.BOTH)

    assert "DNA" in summary
    assert "RNA" in summary
    assert "DNA + RNA" in summary


def test_mode_label_uses_dna_plus_rna_instead_of_both() -> None:
    assert mode_label(MatchMode.DNA) == "DNA only"
    assert mode_label(MatchMode.RNA) == "RNA only"
    assert mode_label(MatchMode.BOTH) == "DNA + RNA"
    assert "Both" not in mode_label(MatchMode.BOTH)


def test_default_rna_subtype_preserves_current_behavior() -> None:
    assert DEFAULT_RNA_SUBTYPE_MODE == RnaSubtypeMode.R_AND_MER


def test_rna_subtype_enabled_only_for_modes_that_use_rna() -> None:
    assert not rna_subtype_enabled(MatchMode.DNA)
    assert rna_subtype_enabled(MatchMode.RNA)
    assert rna_subtype_enabled(MatchMode.BOTH)


def test_matcher_app_rna_subtype_selector_is_compact_and_mode_sensitive(matcher_app) -> None:
    app, root = matcher_app

    assert root.minsize()[0] >= 580
    assert root.minsize()[1] >= 690
    assert root.geometry().startswith("640x690")
    assert app.mode_tiles[MatchMode.BOTH]["title"].cget("text") == "DNA + RNA"
    assert [button.cget("text") for button in app.rna_subtype_buttons.values()] == [
        "R",
        "MeR",
        "R + MeR",
    ]
    assert app.rna_subtype_var.get() == "R + MeR"

    app._select_mode(MatchMode.DNA)
    app._refresh_ui()
    assert {
        button.cget("state")
        for button in app.rna_subtype_buttons.values()
    } == {"disabled"}

    app._select_mode(MatchMode.RNA)
    app._refresh_ui()
    assert {
        button.cget("state")
        for button in app.rna_subtype_buttons.values()
    } == {"normal"}


def test_path_badge_text_distinguishes_default_and_custom() -> None:
    assert path_badge_text(DEFAULT_DNA_PATH, DEFAULT_DNA_PATH) == "Default"
    assert path_badge_text(Path("custom.xlsx"), DEFAULT_DNA_PATH) == "Custom"


def test_database_badge_text_marks_oil_adduct_workbook(tmp_path) -> None:
    path = tmp_path / "oil.xlsx"
    workbook = Workbook()
    workbook.active.title = "Database"
    workbook.create_sheet("DNA 總表")
    workbook.create_sheet("RNA 總表")
    workbook.save(path)

    assert database_badge_text(path, DEFAULT_DNA_PATH) == "Oil adduct"


def test_matcher_app_database_preset_buttons_switch_both_database_paths(matcher_app) -> None:
    app, root = matcher_app

    assert app.standard_database_button is not None
    assert app.oil_adduct_database_button is not None
    assert app.standard_database_button.cget("text") == "Standard DB"
    assert app.oil_adduct_database_button.cget("text") == "Oil Adduct DB"
    assert app.standard_database_button.master is app.oil_adduct_database_button.master
    assert app.standard_database_button.master.master is app.rna_database_group
    assert app.standard_database_button.master.grid_info()["row"] == 2
    assert app.standard_database_button.master.grid_info()["columnspan"] == 2
    assert app.standard_database_button.grid_info()["column"] == 0
    assert app.oil_adduct_database_button.grid_info()["column"] == 1
    assert app.standard_database_button.grid_info()["sticky"] == "ew"
    assert app.oil_adduct_database_button.grid_info()["sticky"] == "ew"
    assert app.standard_database_button.cget("bg") == "#0F766E"
    assert app.oil_adduct_database_button.cget("bg") == "#F0F4F8"

    app._apply_oil_default_database()
    root.update_idletasks()

    assert app.state.dna_db_path == DEFAULT_OIL_ADDUCT_PATH
    assert app.state.rna_db_path == DEFAULT_OIL_ADDUCT_PATH
    assert app.dna_var.get() == str(DEFAULT_OIL_ADDUCT_PATH)
    assert app.rna_var.get() == str(DEFAULT_OIL_ADDUCT_PATH)
    assert app.dna_badge_var.get() == "Oil adduct"
    assert app.rna_badge_var.get() == "Oil adduct"
    assert app.standard_database_button.cget("bg") == "#F0F4F8"
    assert app.oil_adduct_database_button.cget("bg") == "#0F766E"

    app._apply_standard_default_database()
    root.update_idletasks()

    assert app.state.dna_db_path == DEFAULT_DNA_PATH
    assert app.state.rna_db_path == DEFAULT_RNA_PATH
    assert app.dna_var.get() == str(DEFAULT_DNA_PATH)
    assert app.rna_var.get() == str(DEFAULT_RNA_PATH)
    assert app.dna_badge_var.get() == "Default"
    assert app.rna_badge_var.get() == "Default"
    assert app.standard_database_button.cget("bg") == "#0F766E"
    assert app.oil_adduct_database_button.cget("bg") == "#F0F4F8"


def test_matcher_app_custom_database_paths_clear_preset_selection(matcher_app, tmp_path) -> None:
    app, root = matcher_app
    custom_dna = tmp_path / "custom dna database with a very long name.xlsx"
    custom_rna = tmp_path / "custom rna database with a very long name.xlsx"

    app.dna_var.set(str(custom_dna))
    app.rna_var.set(str(custom_rna))
    root.update_idletasks()

    assert app.dna_badge_var.get() == "Custom"
    assert app.rna_badge_var.get() == "Custom"
    assert app.standard_database_button.cget("bg") == "#F0F4F8"
    assert app.oil_adduct_database_button.cget("bg") == "#F0F4F8"


def test_matcher_app_database_preset_row_fits_minimum_window(matcher_app) -> None:
    app, root = matcher_app

    root.deiconify()
    root.geometry("580x650")
    root.update()

    preset_row = app.standard_database_button.master
    assert preset_row.winfo_width() > 0
    assert root.winfo_height() >= 690
    assert app.standard_database_button.winfo_width() >= app.standard_database_button.winfo_reqwidth()
    assert app.standard_database_button.winfo_height() >= app.standard_database_button.winfo_reqheight()
    assert (
        app.oil_adduct_database_button.winfo_width()
        >= app.oil_adduct_database_button.winfo_reqwidth()
    )
    assert (
        app.oil_adduct_database_button.winfo_height()
        >= app.oil_adduct_database_button.winfo_reqheight()
    )
    assert preset_row.winfo_reqwidth() <= app.rna_database_group.winfo_width()
    assert (
        preset_row.winfo_rooty() + preset_row.winfo_height()
        <= app.input_card.winfo_rooty() + app.input_card.winfo_height()
    )


def test_status_appearance_uses_semantic_states() -> None:
    ready = status_appearance("Output folder: C:/demo/Output")
    success = status_appearance("Saved result to: C:/demo/Output/file.xlsx")
    error = status_appearance("Failed: bad dataset")

    assert ready["tone"] == "ready"
    assert success["tone"] == "success"
    assert error["tone"] == "error"


def test_run_matching_in_dna_mode_does_not_require_rna_database(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.csv"
    dna_path = tmp_path / "dna.xlsx"

    pd.DataFrame({"Feature": ["268.1052/17.59(Mz/RT)"]}).to_csv(dataset_path, index=False)
    pd.DataFrame(
        {
            "Short name": ["dX"],
            "Charged monoisotopic mass": [268.1052],
        }
    ).to_excel(dna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=dna_path,
        rna_db_path=tmp_path / "missing-rna.xlsx",
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(state, MatchMode.DNA)
    result = pd.read_excel(output_path)

    assert list(result["Matched Short name"]) == ["dX"]


def test_run_matching_in_dna_mode_accepts_mz_database_column(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    dna_path = tmp_path / "dna.xlsx"

    pd.DataFrame({"Feature": ["268.1052/17.59(Mz/RT)"]}).to_excel(dataset_path, index=False)
    pd.DataFrame({"Short name": ["dX"], "Mz": [268.1052]}).to_excel(dna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=dna_path,
        rna_db_path=tmp_path / "missing-rna.xlsx",
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(state, MatchMode.DNA)
    result = pd.read_excel(output_path)

    assert list(result["Matched Short name"]) == ["dX"]


def test_run_matching_accepts_charged_monoisotopic_mass_in_dataset(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    dna_path = tmp_path / "dna.xlsx"

    pd.DataFrame({"Charged monoisotopic mass": ["268.1052"]}).to_excel(dataset_path, index=False)
    pd.DataFrame({"Short name": ["dX"], "Mz/RT": ["268.1052/9.41"]}).to_excel(dna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=dna_path,
        rna_db_path=tmp_path / "missing-rna.xlsx",
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(state, MatchMode.DNA)
    result = pd.read_excel(output_path)

    assert list(result["Matched Short name"]) == ["dX"]


def test_run_matching_exports_all_feature_sheets_from_excel_workbook(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.xlsx"
    dna_path = tmp_path / "dna.xlsx"

    with pd.ExcelWriter(dataset_path) as writer:
        pd.DataFrame({"Feature": ["268.1052/17.59"]}).to_excel(writer, sheet_name="Summary", index=False)
        pd.DataFrame({"Feature": ["300.0/4.2"]}).to_excel(writer, sheet_name="VIP", index=False)
        pd.DataFrame({"Note": ["skip"]}).to_excel(writer, sheet_name="Metadata", index=False)

    pd.DataFrame(
        {
            "Short name": ["dX", "dY"],
            "Mz": [268.1052, 300.0],
        }
    ).to_excel(dna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=dna_path,
        rna_db_path=tmp_path / "missing-rna.xlsx",
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(state, MatchMode.DNA)
    summary = pd.read_excel(output_path, sheet_name="Summary")
    vip = pd.read_excel(output_path, sheet_name="VIP")
    sheets = pd.ExcelFile(output_path).sheet_names

    assert sheets == ["Summary", "VIP"]
    assert list(summary["Matched Short name"]) == ["dX"]
    assert list(vip["Matched Short name"]) == ["dY"]


def test_run_matching_outputs_formula_column(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.csv"
    dna_path = tmp_path / "dna.xlsx"

    pd.DataFrame({"Feature": ["268.1052/17.59"]}).to_csv(dataset_path, index=False)
    pd.DataFrame({
        "Short name": ["dX"],
        "Charged monoisotopic mass": [268.1052],
        "Formula": ["C10H13N4O4"],
        "Source": ["DNA DB"],
    }).to_excel(dna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=dna_path,
        rna_db_path=tmp_path / "missing-rna.xlsx",
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(state, MatchMode.DNA)
    result = pd.read_excel(output_path)

    assert "Matched Formula" in result.columns
    assert "Matched Source" in result.columns
    assert "Matched Short name" in result.columns
    assert list(result["Matched Formula"]) == ["C10H13N4O4"]
    assert list(result["Matched Source"]) == ["DNA DB"]
    assert list(result["Matched Short name"]) == ["dX"]


def test_run_matching_in_rna_mode_omits_source_column(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.csv"
    rna_path = tmp_path / "rna.xlsx"

    pd.DataFrame({"Feature": ["268.1052/17.59"]}).to_csv(dataset_path, index=False)
    pd.DataFrame(
        {
            "Short name": ["m1A"],
            "[M+H]+": [268.1052],
            "Formula": ["C11H16N5O3"],
        }
    ).to_excel(rna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=tmp_path / "missing-dna.xlsx",
        rna_db_path=rna_path,
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(state, MatchMode.RNA)
    result = pd.read_excel(output_path)

    assert "Matched Source" not in result.columns
    assert list(result["Matched Short name"]) == ["m1A"]


def test_run_matching_in_rna_mode_can_limit_to_mer_only(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.csv"
    rna_path = tmp_path / "rna.xlsx"

    pd.DataFrame({"Feature": ["268.1052/17.59"]}).to_csv(dataset_path, index=False)
    pd.DataFrame(
        {
            "Short name": ["m1A", "m1Am"],
            "[M+H]+": [268.1052, 268.1052],
            "Formula": ["C11H16N5O3", "C11H18N5O4"],
        }
    ).to_excel(rna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=tmp_path / "missing-dna.xlsx",
        rna_db_path=rna_path,
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(
        state,
        MatchMode.RNA,
        rna_subtype_mode=RnaSubtypeMode.MER_ONLY,
    )
    result = pd.read_excel(output_path)

    assert list(result["Matched Short name"]) == ["m1Am"]
    assert list(result["Matched Formula"]) == ["C11H18N5O4"]
    assert "Matched Source" not in result.columns


def test_run_matching_in_dna_mode_ignores_rna_subtype_and_missing_rna_db(tmp_path) -> None:
    dataset_path = tmp_path / "dataset.csv"
    dna_path = tmp_path / "dna.xlsx"

    pd.DataFrame({"Feature": ["268.1052/17.59"]}).to_csv(dataset_path, index=False)
    pd.DataFrame(
        {
            "Short name": ["dX"],
            "Charged monoisotopic mass": [268.1052],
        }
    ).to_excel(dna_path, index=False)

    state = AppState(
        dataset_path=dataset_path,
        dna_db_path=dna_path,
        rna_db_path=tmp_path / "missing-rna.xlsx",
        output_dir=tmp_path / "Output",
    )

    output_path = run_matching(
        state,
        MatchMode.DNA,
        rna_subtype_mode=RnaSubtypeMode.MER_ONLY,
    )
    result = pd.read_excel(output_path)

    assert list(result["Matched Short name"]) == ["dX"]
