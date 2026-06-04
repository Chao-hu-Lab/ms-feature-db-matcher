from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet


DNA_SHEET_NAME = "DNA 總表"
RNA_SHEET_NAME = "RNA 總表"
DATABASE_SHEET_NAME = "Database"

BASE_COLUMNS = [
    "Short name",
    "Name",
    "Formula",
    "Monoisotopic mass",
    "Charged monoisotopic mass",
    "Source",
    "Source Group",
    "Adduct Category",
    "Database Source Sheet",
    "Database Source Row",
]
RNA_COLUMNS = [*BASE_COLUMNS, "RNA Subtype"]

_NO_FILL_KEYS = {"00000000", "00FFFFFF", "FFFFFFFF"}
_SECOND_TABLE_HEADERS = {"category", "chemical", "chemical formula", "exact mass", "m/z", "structure"}


@dataclass(frozen=True)
class OilAdductTables:
    dna: pd.DataFrame
    rna: pd.DataFrame


def is_oil_adduct_workbook(path: Path) -> bool:
    workbook = load_workbook(path, read_only=True)
    return DNA_SHEET_NAME in workbook.sheetnames or RNA_SHEET_NAME in workbook.sheetnames


def build_oil_adduct_tables(path: Path) -> OilAdductTables:
    workbook = load_workbook(path, data_only=True)
    legend = _source_group_legend(workbook)
    dna = _build_table(workbook[DNA_SHEET_NAME], "DNA", legend) if DNA_SHEET_NAME in workbook.sheetnames else pd.DataFrame(columns=BASE_COLUMNS)
    rna = _build_table(workbook[RNA_SHEET_NAME], "RNA", legend) if RNA_SHEET_NAME in workbook.sheetnames else pd.DataFrame(columns=RNA_COLUMNS)
    return OilAdductTables(dna=dna, rna=rna)


def load_oil_adduct_table(path: Path, mode: object) -> pd.DataFrame:
    tables = build_oil_adduct_tables(path)
    mode_text = str(getattr(mode, "value", mode)).upper()
    if mode_text == "DNA":
        return tables.dna
    if mode_text == "RNA":
        return tables.rna
    raise ValueError(f"Unsupported oil adduct database mode: {mode!r}")


def export_oil_adduct_databases(source_path: Path, output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    tables = build_oil_adduct_tables(source_path)
    dna_path = output_dir / "oil_adduct_dna.xlsx"
    rna_path = output_dir / "oil_adduct_rna.xlsx"
    tables.dna.to_excel(dna_path, index=False)
    tables.rna.to_excel(rna_path, index=False)
    return dna_path, rna_path


def _build_table(sheet: Worksheet, mode: str, legend: dict[str, str]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    second_header_row = _find_second_table_header_row(sheet)
    _append_primary_table_rows(rows, sheet, mode, legend, stop_before=second_header_row)
    if second_header_row is not None:
        _append_secondary_table_rows(rows, sheet, mode, second_header_row)
    columns = RNA_COLUMNS if mode == "RNA" else BASE_COLUMNS
    return pd.DataFrame(rows, columns=columns)


def _append_primary_table_rows(
    rows: list[dict[str, object]],
    sheet: Worksheet,
    mode: str,
    legend: dict[str, str],
    *,
    stop_before: int | None,
) -> None:
    current_category = ""
    end_row = stop_before if stop_before is not None else sheet.max_row + 1
    for row_number in range(4, end_row):
        category = _clean_text(sheet.cell(row_number, 2).value)
        if category:
            current_category = category
        name = _clean_text(sheet.cell(row_number, 3).value)
        charged_mass = _coalesce(sheet.cell(row_number, 8).value, sheet.cell(row_number, 6).value)
        if not name or charged_mass == "":
            continue

        source_cell = sheet.cell(row_number, 9)
        source = _clean_text(source_cell.value)
        row = _database_row(
            short_name=_short_name_from_name(name),
            name=name,
            formula=_clean_formula(sheet.cell(row_number, 5).value),
            monoisotopic_mass=_number_or_blank(sheet.cell(row_number, 6).value),
            charged_mass=_number_or_blank(charged_mass),
            source=source,
            source_group=_source_group_from_cell(source_cell, legend, source),
            category=current_category,
            sheet_name=sheet.title,
            row_number=row_number,
        )
        if mode == "RNA":
            row["RNA Subtype"] = _rna_subtype(row["Short name"])
        rows.append(row)


def _append_secondary_table_rows(
    rows: list[dict[str, object]],
    sheet: Worksheet,
    mode: str,
    header_row: int,
) -> None:
    headers = {
        _clean_text(sheet.cell(header_row, column).value).lower(): column
        for column in range(1, sheet.max_column + 1)
        if _clean_text(sheet.cell(header_row, column).value)
    }
    source = _secondary_table_source(sheet, header_row)
    current_category = ""
    for row_number in range(header_row + 1, sheet.max_row + 1):
        category = _clean_text(sheet.cell(row_number, headers["category"]).value)
        if category:
            current_category = category
        name = _clean_text(sheet.cell(row_number, headers["chemical"]).value)
        charged_mass = _number_or_blank(sheet.cell(row_number, headers["m/z"]).value)
        if not name or charged_mass == "":
            continue

        row = _database_row(
            short_name=name,
            name=name,
            formula=_clean_formula(sheet.cell(row_number, headers["chemical formula"]).value),
            monoisotopic_mass=_number_or_blank(sheet.cell(row_number, headers["exact mass"]).value),
            charged_mass=charged_mass,
            source=source,
            source_group=source,
            category=current_category,
            sheet_name=sheet.title,
            row_number=row_number,
        )
        if mode == "RNA":
            row["RNA Subtype"] = _rna_subtype(row["Short name"])
        rows.append(row)


def _database_row(
    *,
    short_name: str,
    name: str,
    formula: str,
    monoisotopic_mass: object,
    charged_mass: object,
    source: str,
    source_group: str,
    category: str,
    sheet_name: str,
    row_number: int,
) -> dict[str, object]:
    return {
        "Short name": short_name,
        "Name": name,
        "Formula": formula,
        "Monoisotopic mass": monoisotopic_mass,
        "Charged monoisotopic mass": charged_mass,
        "Source": source,
        "Source Group": source_group,
        "Adduct Category": category,
        "Database Source Sheet": sheet_name,
        "Database Source Row": row_number,
    }


def _find_second_table_header_row(sheet: Worksheet) -> int | None:
    for row_number in range(1, sheet.max_row + 1):
        values = {_clean_text(sheet.cell(row_number, column).value).lower() for column in range(1, sheet.max_column + 1)}
        if {"category", "chemical", "chemical formula", "exact mass", "m/z"} <= values:
            return row_number
    return None


def _secondary_table_source(sheet: Worksheet, header_row: int) -> str:
    for column in range(1, sheet.max_column + 1):
        value = _clean_text(sheet.cell(header_row, column).value)
        if value and value.lower() not in _SECOND_TABLE_HEADERS:
            return value
    return ""


def _source_group_legend(workbook) -> dict[str, str]:
    if DATABASE_SHEET_NAME not in workbook.sheetnames:
        return {}
    sheet = workbook[DATABASE_SHEET_NAME]
    legend: dict[str, str] = {}
    for row in range(1, min(sheet.max_row, 30) + 1):
        cell = sheet.cell(row, 18)
        key = _fill_key(cell)
        value = _clean_text(cell.value)
        if key and value:
            legend[key] = value
    return legend


def _source_group_from_cell(cell: Cell, legend: dict[str, str], source: str) -> str:
    key = _fill_key(cell)
    if not key:
        return ""
    return legend.get(key, source)


def _fill_key(cell: Cell) -> str:
    fill = cell.fill
    if not fill or not fill.fill_type:
        return ""
    color = fill.fgColor
    if color.type == "rgb":
        key = color.rgb
    elif color.type == "theme":
        key = f"theme:{color.theme}:{color.tint}"
    elif color.type == "indexed":
        key = f"indexed:{color.indexed}"
    else:
        key = f"{color.type}:{color.value}"
    return "" if key in _NO_FILL_KEYS else str(key)


def _short_name_from_name(name: object) -> str:
    text = _clean_text(name)
    if not text.endswith(")"):
        return text
    depth = 0
    for index in range(len(text) - 1, -1, -1):
        char = text[index]
        if char == ")":
            depth += 1
        elif char == "(":
            depth -= 1
            if depth == 0:
                short_name = text[index + 1 : -1].strip()
                return short_name or text
    return text


def _rna_subtype(short_name: object) -> str:
    return "MeR" if "mer" in _clean_text(short_name).lower() else "R"


def _clean_formula(value: object) -> str:
    text = _clean_text(value)
    return text[:-1].strip() if text.endswith("+") else text


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    return "" if text.lower() == "nan" else " ".join(text.split())


def _coalesce(*values: object) -> object:
    for value in values:
        if _clean_text(value):
            return value
    return ""


def _number_or_blank(value: object) -> object:
    text = _clean_text(value)
    if not text:
        return ""
    try:
        return float(text)
    except ValueError:
        return text


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export canonical DNA/RNA oil adduct database workbooks.")
    parser.add_argument("source", type=Path, help="Source oil adduct workbook")
    parser.add_argument("--output-dir", type=Path, default=Path("database"), help="Output directory")
    args = parser.parse_args(argv)

    dna_path, rna_path = export_oil_adduct_databases(args.source, args.output_dir)
    print(dna_path)
    print(rna_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
